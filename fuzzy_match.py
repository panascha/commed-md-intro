#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Fuzzy match nicknames from รับน้อง MD52.xlsx to รายชื่อนักศึกษาCOMMED2569.xlsx
Normalize Thai: remove prefix, vowels, tone marks, diacritics
"""
import openpyxl
import re
from collections import defaultdict

def normalize_thai(s):
    """Remove Thai vowels, tone marks, diacritics, spaces, prefix"""
    if not s:
        return ""
    # Remove prefix
    s = re.sub(r'^(นาย|นางสาว|นาง)\s*', '', str(s).strip())
    # Remove Thai vowels (floating + attached)
    s = re.sub(r'[เแโใไาำะัิีึืุู็่้๊๋์็ๆ]', '', s)
    # Remove spaces
    s = re.sub(r'\s+', '', s)
    return s.lower()

def extract_first_last(name):
    """Split Thai name into first + last"""
    name = re.sub(r'^(นาย|นางสาว|นาง)\s*', '', str(name).strip())
    parts = name.split()
    if len(parts) >= 2:
        return parts[0], parts[-1]
    elif len(parts) == 1:
        return parts[0], ""
    return "", ""

# Load source (nicknames)
src_wb = openpyxl.load_workbook('รับน้อง MD52.xlsx', data_only=True)
nicknames = {}  # normalized_name -> (nickname, original_name, student_id)

for sheet_name in ['น้องสายรหัส', 'RT']:
    ws = src_wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    name_col = hdr.index('ชื่อ') if 'ชื่อ' in hdr else 1
    nick_col = hdr.index('ชื่อเล่น') if 'ชื่อเล่น' in hdr else 4
    id_col = hdr.index('รหัสนักศึกษา') if 'รหัสนักศึกษา' in hdr else 0

    for row in rows[1:]:
        if len(row) > name_col and row[name_col]:
            orig = str(row[name_col]).strip()
            norm = normalize_thai(orig)
            nick = str(row[nick_col]).strip() if len(row) > nick_col and row[nick_col] else ""
            sid = str(row[id_col]).strip() if len(row) > id_col and row[id_col] else ""
            if norm:
                nicknames[norm] = (nick, orig, sid)

# Load target
tgt_wb = openpyxl.load_workbook('รายชื่อนักศึกษาCOMMED2569.xlsx', data_only=True)
ws = tgt_wb['MD 743102']
tgt_rows = list(ws.iter_rows(values_only=True))
header = tgt_rows[3]  # row 4 is header
data_rows = tgt_rows[4:]

# Match
matched = []
unmatched = []

for row in data_rows:
    if not row or not row[2]:  # no student ID
        continue
    sid = str(row[2]).strip()
    name = str(row[3]).strip() if row[3] else ""
    norm = normalize_thai(name)

    # Exact match after normalize
    if norm in nicknames:
        nick, orig, src_id = nicknames[norm]
        matched.append((sid, name, nick, orig, src_id, 'exact'))
        continue

    # Try first name only
    first, last = extract_first_last(name)
    first_norm = normalize_thai(first)

    found = False
    for src_norm, (src_nick, src_orig, src_sid) in nicknames.items():
        src_first, src_last = extract_first_last(src_orig)
        src_first_norm = normalize_thai(src_first)

        # Match first name
        if first_norm and src_first_norm and first_norm == src_first_norm:
            matched.append((sid, name, src_nick, src_orig, src_sid, 'first_name'))
            found = True
            break

        # Fuzzy: last name match
        if last and src_last:
            last_norm = normalize_thai(last)
            src_last_norm = normalize_thai(src_last)
            if last_norm and src_last_norm and last_norm == src_last_norm:
                matched.append((sid, name, src_nick, src_orig, src_sid, 'last_name'))
                found = True
                break

    if not found:
        unmatched.append((sid, name))

# Report
print(f"=== FUZZY MATCH RESULTS ===")
print(f"Total target students: {len(data_rows)}")
print(f"Matched: {len(matched)}")
print(f"Unmatched: {len(unmatched)}")
print()

# Show matched samples
print(f"=== MATCHED SAMPLES (first 20) ===")
for sid, name, nick, orig, src_id, method in matched[:20]:
    print(f"{sid} | {name} → {nick} (from {orig}) [{method}]")
print()

# Show unmatched
print(f"=== UNMATCHED (first 20) ===")
for sid, name in unmatched[:20]:
    print(f"{sid} | {name}")

# Save results
out_wb = openpyxl.Workbook()
ws_out = out_wb.active
ws_out.title = "Matched"
ws_out.append(["รหัส", "ชื่อ", "ชื่อเล่น", "ชื่อต้นฉบับ", "รหัสต้นฉบับ", "วิธี match"])
for row in matched:
    ws_out.append(list(row))

ws_un = out_wb.create_sheet("Unmatched")
ws_un.append(["รหัส", "ชื่อ"])
for row in unmatched:
    ws_un.append(list(row))

out_wb.save('match_results.xlsx')
print(f"\nSaved to match_results.xlsx")
